import os
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from PIL import Image as PILImage

def resolve_image_paths(inspection: Any) -> Tuple[Optional[str], Optional[str]]:
    """
    Locates the actual filesystem paths for original and annotated package images.
    Returns (orig_path, annot_path).
    """
    orig_url = getattr(inspection, "package_image_url", "") or ""
    annot_url = getattr(inspection, "annotated_image_url", "") or ""

    def find_file(url: str) -> Optional[str]:
        if not url:
            return None
        clean = url.lstrip("/").replace("\\", "/")
        candidates = [
            clean,
            os.path.abspath(clean),
            os.path.join(os.getcwd(), clean),
            os.path.join(os.getcwd(), "static", clean.replace("static/", "")),
            os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), clean)
        ]
        for c in candidates:
            if os.path.exists(c) and os.path.isfile(c):
                return os.path.abspath(c)
        return None

    orig_path = find_file(orig_url)
    annot_path = find_file(annot_url)

    # Derive annotated path if not stored explicitly
    if not annot_path and orig_path:
        stem = Path(orig_path).stem
        parent = os.path.dirname(orig_path)
        derived = os.path.join(parent, "annotated", f"{stem}_annotated.jpg")
        if os.path.exists(derived):
            annot_path = derived

    return orig_path, annot_path

def get_rl_image_proportional(img_path: str, max_w: float = 250, max_h: float = 175) -> Optional[RLImage]:
    """Loads image and scales proportionally to fit within max_w and max_h points."""
    if not img_path or not os.path.exists(img_path):
        return None
    try:
        with PILImage.open(img_path) as im:
            orig_w, orig_h = im.size
            if orig_w == 0 or orig_h == 0:
                return None
            aspect = orig_w / orig_h
            calc_w = max_w
            calc_h = calc_w / aspect
            if calc_h > max_h:
                calc_h = max_h
                calc_w = calc_h * aspect
            return RLImage(img_path, width=calc_w, height=calc_h)
    except Exception as e:
        print(f"Error loading image for PDF ({img_path}): {e}")
        return None

def generate_docx_report(inspection: Any, fields: List[Any], violations: List[Any], output_path: str) -> str:
    """
    Generates a fully editable Microsoft Word (.docx) compliance report.
    Embeds BOTH original package photograph and AI-annotated bounding-box image.
    """
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc = Document()

    # Title Header
    title = doc.add_heading("MINISTRY OF CONSUMER AFFAIRS, FOOD & PUBLIC DISTRIBUTION", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("DIRECTORATE OF LEGAL METROLOGY — STATUTORY INSPECTION NOTICE")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.bold = True
    subtitle.runs[0].font.color.rgb = RGBColor(30, 58, 138)

    doc.add_paragraph("Generated under Legal Metrology (Packaged Commodities) Rules, 2011 & LM Act 2009").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("—" * 55).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Inspection Metadata
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_rows = [
        ("Inspection Number:", getattr(inspection, "inspection_number", "N/A")),
        ("Inspection Date:", str(getattr(inspection, "inspection_date", "N/A"))[:19]),
        ("Commodity / Product:", getattr(getattr(inspection, "product", None), "product_name", "Packaged Commodity")),
        ("Manufacturer / Packer:", getattr(getattr(inspection, "company", None), "company_name", "N/A")),
    ]
    for idx, (label, val) in enumerate(meta_rows):
        meta_table.rows[idx].cells[0].paragraphs[0].add_run(label).bold = True
        meta_table.rows[idx].cells[1].paragraphs[0].add_run(str(val))

    doc.add_paragraph()

    # Overall Verdict Callout
    status = getattr(inspection, "overall_status", "COMPLIANT")
    verdict_para = doc.add_paragraph()
    verdict_para.add_run(f"COMPLIANCE VERDICT: {status.replace('_', ' ')} (Score: {getattr(inspection, 'compliance_score', 100)}%)\n").bold = True

    # 1. Statutory Declarations Assessment (Legal Metrology & Packaged Food)
    lm_fields = [f for f in fields if "[LM]" in getattr(f, "field_name", "")]
    food_fields = [f for f in fields if "[FOOD]" in getattr(f, "field_name", "")]
    general_fields = [f for f in fields if "[LM]" not in getattr(f, "field_name", "") and "[FOOD]" not in getattr(f, "field_name", "")]

    def add_docx_table(title_text, field_items, header_color="1E3A8A"):
        doc.add_heading(title_text, level=2)
        dec_table = doc.add_table(rows=1, cols=4)
        dec_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["Declaration Field", "Extracted Value", "Confidence", "Status"]
        for i, h in enumerate(headers):
            cell = dec_table.rows[0].cells[i]
            run = cell.paragraphs[0].add_run(h)
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            tcPr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), header_color)
            tcPr.append(shd)

        for f in field_items:
            row = dec_table.add_row()
            fname = getattr(f, "field_name", "").replace("[LM] ", "").replace("[FOOD] ", "")
            row.cells[0].paragraphs[0].text = fname
            row.cells[1].paragraphs[0].text = getattr(f, "extracted_value", "") or "[NOT DETECTED]"
            row.cells[2].paragraphs[0].text = f"{int(getattr(f, 'confidence_score', 95))}%"
            row.cells[3].paragraphs[0].text = getattr(f, "field_status", "Verified")

        doc.add_paragraph()

    if lm_fields:
        add_docx_table("1. Mandatory Legal Metrology Declarations (Packaged Commodities Rules, 2011)", lm_fields, "1E3A8A")
    if food_fields:
        add_docx_table("2. Packaged Food Safety & Standards Assessment (FSSAI Regulations)", food_fields, "047857")
    if general_fields and not lm_fields and not food_fields:
        add_docx_table("1. Mandatory Packaging Declarations Assessment (Rule 6)", general_fields, "1E3A8A")

    # 2. Photographic Evidence Section (Both Original and Annotated Images)
    orig_file, annot_file = resolve_image_paths(inspection)
    if orig_file or annot_file:
        doc.add_heading("2. Photographic Evidence & Machine Bounding Box Inspection Map", level=2)
        doc.add_paragraph("Statutory package image captured during inspection alongside OpenCV bounding box detection evidence map:")

        if orig_file and annot_file:
            img_tbl = doc.add_table(rows=2, cols=2)
            img_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

            p0 = img_tbl.rows[0].cells[0].paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p0.add_run().add_picture(orig_file, width=Inches(2.85))

            p1 = img_tbl.rows[0].cells[1].paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p1.add_run().add_picture(annot_file, width=Inches(2.85))

            c0 = img_tbl.rows[1].cells[0].paragraphs[0]
            c0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = c0.add_run("Figure 1: Captured Package Photograph")
            r0.font.size = Pt(8.5)
            r0.font.italic = True
            r0.font.color.rgb = RGBColor(100, 116, 139)

            c1 = img_tbl.rows[1].cells[1].paragraphs[0]
            c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r1 = c1.add_run("Figure 2: AI Readability & Bounding Box Map")
            r1.font.size = Pt(8.5)
            r1.font.italic = True
            r1.font.color.rgb = RGBColor(100, 116, 139)

        elif orig_file:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(orig_file, width=Inches(4.2))
            cap = doc.add_paragraph("Figure 1: Captured Package Photograph")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.runs[0].font.size = Pt(8.5)
            cap.runs[0].font.italic = True

        doc.add_paragraph()

    # 3. Violations Section (if any)
    if violations:
        doc.add_heading("3. Statutory Rule Infringements & Violations Detected", level=2)
        v_table = doc.add_table(rows=1, cols=4)
        v_headers = ["Rule Citation", "Declaration Field", "Severity", "Statutory Penalty"]
        for i, h in enumerate(v_headers):
            cell = v_table.rows[0].cells[i]
            r = cell.paragraphs[0].add_run(h)
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            tcPr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), 'DC2626')
            tcPr.append(shd)

        for v in violations:
            row = v_table.add_row()
            row.cells[0].paragraphs[0].text = getattr(v, "rule_name", "")
            row.cells[1].paragraphs[0].text = getattr(v, "field_name", "")
            row.cells[2].paragraphs[0].text = getattr(v, "severity", "Medium")
            row.cells[3].paragraphs[0].text = f"Rs. {getattr(v, 'penalty_amount', 10000.0):,.2f}"

        doc.add_paragraph()

    # 4. Inspector Decision & Editable Remarks
    doc.add_heading("4. Enforcement Officer Review & Decision", level=2)
    doc.add_paragraph(f"Decision: {str(getattr(inspection, 'inspector_decision', 'Under Review')).upper()}").runs[0].font.bold = True
    doc.add_paragraph(f"Official Remarks / Directions:\n{getattr(inspection, 'inspector_remarks', 'Verified in routine inspection. No compoundable action required.')}")

    doc.add_paragraph("\n\n" + "_" * 40)
    doc.add_paragraph("Inspector Signature & Official Stamp\nDirectorate of Legal Metrology")

    doc.save(output_path)
    return output_path

def generate_pdf_report(inspection: Any, fields: List[Any], violations: List[Any], output_path: str) -> str:
    """
    Generates an official sealed PDF compliance report using ReportLab.
    Embeds BOTH original package photograph and AI-annotated bounding-box image.
    """
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc = SimpleDocTemplate(output_path, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    header_style = ParagraphStyle(
        'GovHeader',
        parent=styles['Heading1'],
        fontSize=13,
        leading=16,
        alignment=1,
        textColor=colors.HexColor('#1E3A8A')
    )
    sub_style = ParagraphStyle(
        'GovSub',
        parent=styles['Normal'],
        fontSize=9,
        leading=12,
        alignment=1,
        textColor=colors.HexColor('#475569')
    )
    caption_style = ParagraphStyle(
        'ImgCaption',
        parent=styles['Normal'],
        fontSize=8,
        leading=10,
        alignment=1,
        textColor=colors.HexColor('#475569')
    )

    story = []

    # Letterhead Header
    story.append(Paragraph("MINISTRY OF CONSUMER AFFAIRS, FOOD & PUBLIC DISTRIBUTION", header_style))
    story.append(Paragraph("DEPARTMENT OF CONSUMER AFFAIRS — DIRECTORATE OF LEGAL METROLOGY", sub_style))
    story.append(Paragraph("Statutory Package Compliance Audit Report (Rules 2011)", sub_style))
    story.append(Spacer(1, 14))

    # Inspection Metadata Table
    status_str = getattr(inspection, 'overall_status', 'COMPLIANT').replace('_', ' ')
    status_color = colors.HexColor('#16A34A') if 'COMPLIANT' in status_str else colors.HexColor('#DC2626')

    meta_data = [
        [Paragraph("<b>Inspection ID:</b>", styles['Normal']), Paragraph(getattr(inspection, 'inspection_number', 'N/A'), styles['Normal']),
         Paragraph("<b>Date:</b>", styles['Normal']), Paragraph(str(getattr(inspection, 'inspection_date', 'N/A'))[:19], styles['Normal'])],
        [Paragraph("<b>Commodity:</b>", styles['Normal']), Paragraph(getattr(getattr(inspection, 'product', None), 'product_name', 'Commodity'), styles['Normal']),
         Paragraph("<b>Score:</b>", styles['Normal']), Paragraph(f"{getattr(inspection, 'compliance_score', 100)}%", styles['Normal'])],
        [Paragraph("<b>Manufacturer:</b>", styles['Normal']), Paragraph(getattr(getattr(inspection, 'company', None), 'company_name', 'N/A'), styles['Normal']),
         Paragraph("<b>Status:</b>", styles['Normal']), Paragraph(f"<font color='{status_color.hexval()}'><b>{status_str}</b></font>", styles['Normal'])],
    ]
    meta_table = Table(meta_data, colWidths=[90, 180, 70, 180])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 12))

    # 1. Declarations Tables (Legal Metrology & Packaged Food)
    lm_fields = [f for f in fields if "[LM]" in getattr(f, "field_name", "")]
    food_fields = [f for f in fields if "[FOOD]" in getattr(f, "field_name", "")]
    general_fields = [f for f in fields if "[LM]" not in getattr(f, "field_name", "") and "[FOOD]" not in getattr(f, "field_name", "")]

    def add_pdf_table(title, field_items, header_hex="#1E3A8A"):
        story.append(Paragraph(f"<b>{title}</b>", styles['Heading3']))
        table_data = [["Field", "Extracted Value", "Confidence", "Status"]]
        for f in field_items:
            fname = getattr(f, "field_name", "").replace("[LM] ", "").replace("[FOOD] ", "")
            table_data.append([
                fname,
                getattr(f, "extracted_value", "")[:50] or "[NOT DETECTED]",
                f"{int(getattr(f, 'confidence_score', 95))}%",
                getattr(f, "field_status", "Verified")
            ])
        tbl = Table(table_data, colWidths=[140, 240, 70, 70])
        tbl.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(header_hex)),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3.5),
            ('TOPPADDING', (0, 0), (-1, -1), 3.5),
        ]))
        story.append(tbl)
        story.append(Spacer(1, 10))

    if lm_fields:
        add_pdf_table("1. Mandatory Legal Metrology Declarations (Packaged Commodities Rules, 2011)", lm_fields, "#1E3A8A")
    if food_fields:
        add_pdf_table("2. Packaged Food Safety & Standards Assessment (FSSAI Regulations)", food_fields, "#047857")
    if general_fields and not lm_fields and not food_fields:
        add_pdf_table("1. Mandatory Declarations (Legal Metrology Rule 6)", general_fields, "#1E3A8A")

    # 2. Photographic Evidence (Dual-Image Comparison Table)
    orig_file, annot_file = resolve_image_paths(inspection)
    if orig_file or annot_file:
        story.append(Paragraph("<b>2. Photographic Evidence & Machine Bounding Box Inspection Map</b>", styles['Heading3']))
        if orig_file and annot_file:
            rl_orig = get_rl_image_proportional(orig_file, max_w=255, max_h=175)
            rl_annot = get_rl_image_proportional(annot_file, max_w=255, max_h=175)
            if rl_orig and rl_annot:
                img_data = [
                    [rl_orig, rl_annot],
                    [Paragraph("<b>Figure 1:</b> Captured Package Photograph", caption_style),
                     Paragraph("<b>Figure 2:</b> AI Bounding-Box Inspection Map", caption_style)]
                ]
                img_table = Table(img_data, colWidths=[270, 270])
                img_table.setStyle(TableStyle([
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                    ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
                    ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
                    ('TOPPADDING', (0, 0), (-1, -1), 6),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                ]))
                story.append(img_table)
                story.append(Spacer(1, 12))
        elif orig_file:
            rl_orig = get_rl_image_proportional(orig_file, max_w=380, max_h=200)
            if rl_orig:
                story.append(rl_orig)
                story.append(Paragraph("<b>Figure 1:</b> Captured Package Photograph", caption_style))
                story.append(Spacer(1, 12))

    # 3. Violations (if any)
    if violations:
        story.append(Paragraph("<b>3. Statutory Non-Compliances & Infringements</b>", styles['Heading3']))
        v_data = [["Rule Citation", "Declaration", "Severity", "Penalty"]]
        for v in violations:
            v_data.append([
                getattr(v, "rule_name", "")[:40],
                getattr(v, "field_name", ""),
                getattr(v, "severity", "High"),
                f"Rs. {getattr(v, 'penalty_amount', 10000.0):,.0f}"
            ])
        v_table = Table(v_data, colWidths=[180, 140, 80, 120])
        v_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#DC2626')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(v_table)
        story.append(Spacer(1, 12))

    # 4. Officer Remarks & Stamp
    story.append(Paragraph(f"<b>4. Officer Decision:</b> {getattr(inspection, 'inspector_decision', 'Confirmed') or 'Confirmed'}", styles['Normal']))
    story.append(Paragraph(f"<b>Official Remarks:</b> {getattr(inspection, 'inspector_remarks', 'Verified in routine inspection.') or 'Verified in routine inspection.'}", styles['Normal']))
    story.append(Spacer(1, 20))
    story.append(Paragraph("____________________________________________<br/>Authorised Legal Metrology Officer Signature & Seal", styles['Normal']))

    doc.build(story)
    return output_path
